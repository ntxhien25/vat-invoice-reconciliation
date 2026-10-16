import sys
from pathlib import Path
sys.path.insert(0,str(Path(__file__).resolve().parents[1]))
"""Reconciliation tests using synthetic records, not company invoices."""
import unittest
import tempfile
from pathlib import Path
import pandas as pd
from vat_reconciliation import (SCHEMA, finish, identifier, reconcile, status_summary,
                                find_export, misa_record_status, source_review_views,
                                matching_overview, source_status_overview, plot_comparison)


def frame(source='gdt', number='000123', vat=80, imported=False, date='01/08/2026', status='Hóa đơn mới'):
    data = finish([['purchases', '0123456789', identifier(number, True),
                    'C26TXX' if source == 'gdt' else '', date, 1000, vat, status,
                    imported, source + '.xlsx', 'Invoices', 7]])
    data['total_amount'] = 1000 + (vat if vat is not None else 0)
    return data


class ReconciliationTests(unittest.TestCase):
    def test_display_formatting_preserves_zero_identifiers_and_source_data(self):
        from vat_reconciliation import format_table_value, display_records
        self.assertEqual(format_table_value(0), '-')
        self.assertEqual(format_table_value(0.0), '-')
        self.assertEqual(format_table_value('0000'), '0000')
        self.assertEqual(format_table_value(77451360), '77,451,360')
        self.assertIs(format_table_value(False), False)
        source = frame(vat=0)
        view = display_records(source)
        self.assertIn("Seller's Tax Code", view)
        self.assertEqual(view['Source Status'].iloc[0], 'New Invoice')
        self.assertEqual(source.status.iloc[0], 'Hóa đơn mới')
        self.assertEqual(source.vat_amount.iloc[0], 0)
        source.loc[0, 'direction'] = 'sales'
        self.assertIn("Purchaser's Tax Code", display_records(source))

    def test_reader_accepts_titled_and_legacy_names(self):
        with tempfile.TemporaryDirectory() as folder:
            p = Path(folder) / 'Purchases_coded_2026-08.xlsx'
            p.touch()
            self.assertEqual(find_export(folder, 'purchases_coded_2026-08.xlsx'), p)
            p.unlink()
            p = Path(folder) / 'purchases_coded_2026-08.xlsx'
            p.touch()
            self.assertEqual(find_export(folder, 'Purchases_coded_2026-08.xlsx'), p)

    def test_blank_notes_are_new_and_original_adjustment_note_retained(self):
        self.assertEqual(misa_record_status('', False), 'Hóa đơn mới')
        self.assertEqual(misa_record_status('Hóa đơn đã bị điều chỉnh', False), 'Hóa đơn đã bị điều chỉnh')
        self.assertEqual(misa_record_status('', True, False), 'Customs Declaration')
        self.assertEqual(misa_record_status('', True, True), 'Hóa đơn mới')
        self.assertEqual(misa_record_status('Special case: adjustment reference', False, direction='purchases'), '')
        self.assertEqual(misa_record_status('', True), 'Hóa đơn mới')

    def test_twelve_digit_misa_input_document_review_label(self):
        misa = frame('misa', number='980000000001', imported=True)
        misa.loc[0, 'party_tax_id'] = ''
        result = reconcile(finish([]), misa, '2026-08', 'purchases')
        self.assertEqual(source_review_views(result)['MISA']['Review Reason'].tolist(), ['Custom Declaration'])
        self.assertEqual(result.result.tolist(), ['missing_matching_fields'])
        misa.loc[0, 'direction'] = 'sales'
        result = reconcile(finish([]), misa, '2026-08', 'sales')
        self.assertEqual(source_review_views(result)['MISA']['Review Reason'].tolist(), ['Missing Matching Fields'])

    def test_each_source_review_count_balances_with_matching_candidates(self):
        gdt = pd.concat([frame(), frame(number='99')], ignore_index=True)
        misa = frame('misa')
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        reviews = source_review_views(result)
        self.assertEqual(len(reviews['GDT']), 1)
        self.assertEqual(len(reviews['MISA']), 0)
        summary = matching_overview({'purchases': result}, gdt, misa)
        self.assertEqual(summary['Records Requiring Review'].tolist(), [1, 0])
        self.assertEqual(summary['Matched Invoices'].tolist(), [1, 1])

    def test_status_summary_totals_and_chart(self):
        import matplotlib
        matplotlib.use('Agg')
        import matplotlib.pyplot as plt
        summary = source_status_overview(frame(), frame('misa'))
        for _, row in summary.iterrows():
            self.assertEqual(row['Total Invoices'], sum(row[c] for c in ['New Invoices', 'Adjustment Invoices', 'Replacement Invoices', 'Cancelled Invoices', 'Other']))
        fig = plot_comparison(summary, '2026-08')
        self.assertEqual(len(fig.axes), 2)
        plt.close(fig)

    def test_padding_is_normalised_only_in_comparison_keys(self):
        self.assertEqual(identifier('000123', True), '000123')
        self.assertEqual(identifier('0123456789'), '0123456789')
        result = reconcile(frame(), frame('misa', '123'), '2026-08', 'purchases')
        self.assertEqual(result.result.tolist(), ['matched_amounts_agree'])
        self.assertEqual(result.invoice_number_gdt.iloc[0], '000123')
        self.assertEqual(result.invoice_number_misa.iloc[0], '123')
        self.assertEqual(result.confidence.iloc[0], 'matched_without_misa_series')

    def test_tax_code_padding_is_normalised_without_changing_source_values(self):
        gdt, misa = frame(), frame('misa', '123')
        misa.loc[0, 'party_tax_id'] = '123456789'
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        self.assertEqual(result.result.tolist(), ['matched_amounts_agree'])
        self.assertEqual(result.party_tax_id_gdt.iloc[0], '0123456789')
        self.assertEqual(result.party_tax_id_misa.iloc[0], '123456789')
        self.assertEqual(gdt.party_tax_id.iloc[0], '0123456789')

    def test_four_sheet_export_preserves_identifiers_and_partitions_reviews(self):
        from openpyxl import load_workbook
        from vat_reconciliation import reconciliation_export_tables, export_reconciliation_workbook
        gdt = pd.concat([frame(), frame(number='000099'), frame(number='55', vat=60)], ignore_index=True)
        misa = pd.concat([frame('misa'), frame('misa', number='000088'), frame('misa', number='55'), frame('misa', number='33', imported=True)], ignore_index=True)
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        tables = reconciliation_export_tables(result)
        self.assertEqual([len(t) for t in tables.values()], [1, 1, 1, 2])
        with tempfile.TemporaryDirectory() as directory:
            path = export_reconciliation_workbook(result, Path(directory) / 'review.xlsx', source_layouts=False)
            book = load_workbook(path)
            self.assertEqual(book.sheetnames, list(tables))
            sheet = book.worksheets[0]
            headers = [c.value for c in sheet[1]]
            for header, expected in [('GDT Invoice Number', '000123'), ('MISA Invoice Number', '000123'), ("GDT Seller's Tax Code", '0123456789')]:
                cell = sheet.cell(2, headers.index(header) + 1)
                self.assertEqual(cell.value, expected)
                self.assertEqual(cell.data_type, 's')
                self.assertEqual(cell.number_format, '@')
            self.assertEqual(sheet.cell(2, headers.index('GDT Amount Excluding VAT') + 1).number_format, '#,##0;[Red](#,##0);"-"')
            book.close()

    def test_original_layout_export_consolidates_matches_and_stacks_mismatches(self):
        from openpyxl import Workbook, load_workbook
        from vat_reconciliation import export_reconciliation_workbook, original_header_export_tables
        with tempfile.TemporaryDirectory() as directory:
            directory = Path(directory)
            gdt_path, misa_path = directory / 'Purchases_coded.xlsx', directory / 'misa.xlsx'
            g_headers = ['STT', 'Ký hiệu hóa đơn', 'Số hóa đơn', 'Ngày lập', 'MST người bán', 'Tên người bán', 'Tổng tiền chưa thuế', 'Tổng tiền thuế', 'Tổng tiền thanh toán']
            m_headers = ['STT', 'Số hóa đơn', 'Ngày, tháng, năm lập hóa đơn', 'Mã số thuế người bán', 'Tên người bán', 'Giá trị HHDV mua vào chưa có thuế', 'Thuế GTGT đủ điều kiện khấu trừ thuế', 'Ghi chú']
            for path, rows in [(gdt_path, [g_headers, [1, 'C26TXX', '123', '01/08/2026', '0123456789', 'Supplier', 1000, 80, 1080], [2, 'C26TXX', '55', '01/08/2026', '0123456789', 'Supplier', 1000, 100, 1100]]),
                               (misa_path, [m_headers, [1, '000123', '01/08/2026', '0123456789', 'Supplier', 1000, 80, ''], [2, '000055', '01/08/2026', '0123456789', 'Supplier', 1000, 80, '']])]:
                book = Workbook();sheet = book.active;sheet.title = 'Invoices'
                for row in rows:sheet.append(row)
                book.save(path);book.close()
            gdt = pd.concat([frame(number='123'), frame(number='55', vat=100)], ignore_index=True)
            misa = pd.concat([frame('misa'), frame('misa', number='000055')], ignore_index=True)
            for data, path in [(gdt, gdt_path), (misa, misa_path)]:
                data['source_file'] = str(path);data['source_row'] = [2, 3]
            result = reconcile(gdt, misa, '2026-08', 'purchases')
            tables = original_header_export_tables(result)
            matched = tables['Matching Reconciliations']
            self.assertNotIn('Result', matched)
            self.assertNotIn('GDT Số hóa đơn', matched)
            self.assertNotIn('MISA Số hóa đơn', matched)
            self.assertEqual(matched['Invoice Number'].tolist(), ['000123'])
            self.assertEqual(matched['Invoice Total'].tolist(), [1080])
            self.assertEqual(len(tables['Mismatch Reconciliations'].attrs['vertical_blocks']), 1)
            exported = export_reconciliation_workbook(result, directory / 'review.xlsx')
            book = load_workbook(exported)
            self.assertEqual(book.worksheets[0].cell(1, 1).value, 'Invoice Number')
            self.assertEqual(book.worksheets[0].cell(2, 1).value, '000123')
            sheet = book['Mismatch Reconciliations']
            gdt_heading = next(row[0].row for row in sheet if row[0].value == 'GDT')
            misa_heading = next(row[0].row for row in sheet if row[0].value == 'MISA')
            self.assertLess(gdt_heading, misa_heading)
            self.assertEqual([sheet.cell(gdt_heading+2, i+1).value for i in range(len(g_headers))], g_headers)
            self.assertEqual(sheet.cell(misa_heading+2, g_headers.index('Số hóa đơn')+1).value, 'Số hóa đơn')
            self.assertEqual(sheet.cell(gdt_heading+3, 3).value, '55')
            self.assertEqual(sheet.cell(misa_heading+3, g_headers.index('Số hóa đơn')+1).value, '000055')
            book.close()

    def test_deductible_vat_difference_is_review_not_error(self):
        result = reconcile(frame(), frame('misa', vat=0), '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'amount_difference_requires_review')
        self.assertEqual(result.vat_difference.iloc[0], -80)

    def test_blank_gdt_vat_compares_as_zero_without_changing_source(self):
        result = reconcile(frame(vat=None), frame('misa', vat=0), '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'matched_amounts_agree')
        self.assertTrue(pd.isna(result.vat_amount_gdt.iloc[0]))
        self.assertEqual(result.vat_difference.iloc[0], 0)

    def test_blank_gdt_net_is_commercial_review_even_when_total_and_vat_match(self):
        gdt, misa = frame(vat=None), frame('misa', vat=0)
        gdt.loc[0, 'net_amount'] = None
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'commercial_invoice_requires_review')
        self.assertEqual(result.total_difference.iloc[0], 0)
        self.assertTrue(pd.isna(gdt.net_amount.iloc[0]))

        reviews = source_review_views(result)
        self.assertEqual(reviews['GDT']['Review Reason'].tolist(), ['Commercial Invoices'])
        self.assertEqual(reviews['MISA']['Review Reason'].tolist(), ['Commercial Invoices'])
        self.assertTrue(pd.isna(result.net_difference.iloc[0]))
        summary = matching_overview({'purchases': result}, gdt, misa)
        self.assertEqual(summary['Invoices Matched by Key'].tolist(), [1, 1])
        self.assertEqual(summary['Matched Invoices'].tolist(), [0, 0])

    def test_net_difference_is_reviewed_when_total_and_vat_match(self):
        gdt, misa = frame(), frame('misa')
        gdt.loc[0, 'net_amount'] = 900
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'amount_difference_requires_review')
        self.assertEqual(result.net_difference.iloc[0], 100)
        self.assertEqual(result.total_difference.iloc[0], 0)

    def test_payment_total_difference_is_detected_with_equal_net_and_vat(self):
        gdt, misa = frame(), frame('misa')
        gdt.loc[0, 'total_amount'] += 500
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'amount_difference_requires_review')
        self.assertEqual(result.total_difference.iloc[0], -500)

    def test_blank_gdt_total_compares_as_zero_and_misa_missing_remains_missing(self):
        gdt, misa = frame(), frame('misa')
        gdt.loc[0, 'total_amount'] = None
        result = reconcile(gdt, misa, '2026-08', 'purchases')
        self.assertEqual(result.total_difference.iloc[0], 1080)
        misa.loc[0, 'total_amount'] = None
        self.assertEqual(reconcile(gdt, misa, '2026-08', 'purchases').result.iloc[0], 'missing_amount_requires_review')

    def test_duplicate_key_does_not_create_cartesian_matches(self):
        gdt = pd.concat([frame(), frame()], ignore_index=True)
        result = reconcile(gdt, frame('misa'), '2026-08', 'purchases')
        self.assertEqual(len(result), 3)
        self.assertEqual(result.result.value_counts().to_dict(),
                         {'ambiguous_matching_key': 2, 'counterpart_requires_review': 1})

    def test_imports_and_prior_period_documents_are_retained(self):
        misa = pd.concat([frame('misa', imported=True), frame('misa', number='999', date='29/07/2026')], ignore_index=True)
        result = reconcile(finish([]), misa, '2026-08', 'purchases')
        self.assertEqual(result.result.tolist(), ['misa_only', 'misa_only'])
        self.assertEqual(int(result.import_document_misa.sum()), 1)

    def test_cancelled_invoice_is_paired_by_key_with_original_status_retained(self):
        result = reconcile(frame(status='Hóa đơn đã hủy'), frame('misa'), '2026-08', 'purchases')
        self.assertEqual(result.result.tolist(), ['matched_amounts_agree'])
        self.assertEqual(result.status_gdt.iloc[0], 'Hóa đơn đã hủy')

    def test_statuses_partition_with_unknown_and_overlap_flags(self):
        data = pd.concat([frame(status=s) for s in ['Hóa đơn mới', 'Hóa đơn điều chỉnh thay thế', 'Trạng thái mới chưa biết']], ignore_index=True)
        summary, detail = status_summary(data)
        self.assertEqual(summary['count'].sum(), 3)
        self.assertEqual(int(detail.overlapping_labels.sum()), 1)
        self.assertEqual(summary.set_index('category').loc['other', 'count'], 1)

    def test_missing_counterparty_is_not_matched(self):
        gdt = frame()
        gdt.loc[0, 'party_tax_id'] = ''
        result = reconcile(gdt, finish([]), '2026-08', 'purchases')
        self.assertEqual(result.result.iloc[0], 'missing_matching_fields')


if __name__ == '__main__':
    unittest.main(verbosity=2)
