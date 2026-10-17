"""Reconcile GDT invoice exports with MISA VAT return schedules.

Match purchase invoices by invoice number, issue date and seller's tax code;
match sales invoices using the purchaser's tax code. Compare invoice totals and
VAT amounts, retaining source references for follow-up. MISA input VAT is the
amount recorded as deductible, which may differ from VAT shown on an invoice.
The MISA schedules omit invoice series, so duplicate matching keys need review.

Vietnamese column names in the readers match the original source files.
Presentation labels and project explanations use English.
"""
from __future__ import annotations
import re
import unicodedata
import warnings
from pathlib import Path
import pandas as pd
from openpyxl import load_workbook

SCHEMA = ['direction', 'party_tax_id', 'invoice_number', 'series', 'issue_date',
          'net_amount', 'vat_amount', 'status', 'import_document', 'source_file',
          'source_sheet', 'source_row']

STATUS_LABELS = {'new': 'New Invoices', 'adjustment_related': 'Adjustment Invoices',
                 'replacement_related': 'Replacement Invoices', 'cancelled': 'Cancelled Invoices', 'other': 'Other'}

PAIRED_RESULTS = ('matched_amounts_agree', 'amount_difference_requires_review',
                  'missing_amount_requires_review', 'commercial_invoice_requires_review')


def text(value):
    return '' if value is None or pd.isna(value) else re.sub(r'\s+', ' ', str(value)).strip()


def identifier(value, invoice=False):
    """Retain source identifiers, including leading zeros; never infer padding."""
    if value is None or pd.isna(value):
        return ''
    if isinstance(value, float) and value.is_integer():
        return str(int(value))  # Numeric Excel cells contain no stored leading zeros.
    return str(value)


def folded(value):
    return ''.join(c for c in unicodedata.normalize('NFD', text(value))
                   if unicodedata.category(c) != 'Mn').upper().replace('Đ', 'D')


def read_book(path):
    path = Path(path)
    if not path.is_file():
        raise ValueError(f'Missing file: {path}')
    # Suppress only the known exporter-style warning, not all openpyxl warnings.
    with warnings.catch_warnings():
        warnings.filterwarnings('ignore', message='Workbook contains no default style.*', category=UserWarning)
        book = load_workbook(path, read_only=True, data_only=True)
        try:
            result = {}
            for sheet in book:
                sheet.reset_dimensions()  # GDT's declared A1 range can be wrong.
                rows = list(sheet.iter_rows(values_only=True))
                width = max((len(row) for row in rows), default=0)
                result[sheet.title] = [tuple(row) + (None,) * (width - len(row)) for row in rows]
            return result
        finally:
            book.close()


def finish(records):
    frame = pd.DataFrame(records, columns=SCHEMA)
    frame['issue_date'] = pd.to_datetime(frame['issue_date'], dayfirst=True, errors='coerce')
    for column in ('net_amount', 'vat_amount'):
        original = frame[column]
        frame[column] = pd.to_numeric(original, errors='coerce')
        if (original.map(text).ne('') & frame[column].isna()).any():
            raise ValueError(f'Unrecognised numeric values in {column}; source needs review.')
    return frame


def read_gdt(path, direction, tax_id, month, sheet_names=None):
    records, totals, found = [], [], False
    for sheet, rows in read_book(path).items():
        if sheet_names and sheet not in sheet_names:
            continue
        for index, row in enumerate(rows[:30]):
            headers = [text(v) for v in row]
            if {'Số hóa đơn', 'Ký hiệu hóa đơn', 'Ngày lập'}.issubset(headers):
                break
        else:
            continue
        found = True
        def col(prefix):
            indices = [i for i, h in enumerate(headers) if h.startswith(prefix)]
            if len(indices) != 1:
                raise ValueError(f'{Path(path).name}/{sheet}: ambiguous or missing column {prefix}')
            return indices[0]
        number, series, issued = col('Số hóa đơn'), col('Ký hiệu hóa đơn'), col('Ngày lập')
        own = col('MST người mua' if direction == 'purchases' else 'MST người bán')
        party = col('MST người bán' if direction == 'purchases' else 'MST người mua')
        net = col('Tổng tiền chưa thuế') if 'Tổng tiền chưa thuế' in headers else None
        vat = col('Tổng tiền thuế') if 'Tổng tiền thuế' in headers else None
        total = col('Tổng tiền thanh toán')
        status = col('Trạng thái hóa đơn') if 'Trạng thái hóa đơn' in headers else None
        for source_row, row in enumerate(rows[index + 1:], index + 2):
            if not text(row[number]):
                continue
            if text(row[number]) == 'Số hóa đơn':
                continue
            if identifier(row[own]) != tax_id:
                raise ValueError(f'{Path(path).name}: taxpayer does not match {tax_id}')
            records.append([direction, identifier(row[party]), identifier(row[number], True),
                            identifier(row[series]), row[issued], row[net] if net is not None else None,
                            row[vat] if vat is not None else None, text(row[status]) if status is not None else '',
                            False, str(Path(path).resolve()), sheet, source_row])
            totals.append(row[total])
    if not found:
        raise ValueError(f'{Path(path).name}: blank workbook or unrecognised GDT invoice table.')
    frame = finish(records)
    frame['total_amount'] = pd.to_numeric(pd.Series(totals, index=frame.index), errors='coerce')
    if any(text(raw) and pd.isna(parsed) for raw, parsed in zip(totals, frame.total_amount)):
        raise ValueError(f'{Path(path).name}: unrecognised total payment amount.')
    if frame['issue_date'].isna().any() or frame['issue_date'].dt.strftime('%Y-%m').ne(month).any():
        raise ValueError(f'{Path(path).name}: missing dates or invoices outside {month}')
    return frame


def misa_record_status(notes, import_document, customs_as_new=True, direction='sales'):
    """Classify blank sales notes and customs declarations; retain purchase special-case notes."""
    if import_document:
        return 'Hóa đơn mới' if customs_as_new else 'Customs Declaration'
    if direction == 'purchases' and text(notes):
        return ''  # A purchase special-case note is not evidence of invoice status.
    return text(notes) or 'Hóa đơn mới'


def read_misa(path, month, *, customs_as_new=True):
    records, extra, found_directions, declared_month = [], [], set(), None
    for sheet, rows in read_book(path).items():
        intro = ' '.join(text(v) for row in rows[:7] for v in row)
        period = re.search(r'Tháng\s+(\d+)\s+năm\s+(\d{4})', intro, re.I)
        if period:
            declared_month = f'{period[2]}-{int(period[1]):02d}'
        title = folded(intro)
        if 'BANG KE HOA DON' not in title:
            continue
        direction = 'purchases' if 'MUA VAO' in title else 'sales' if 'BAN RA' in title else None
        if not direction:
            continue
        found_directions.add(direction)
        for index, row in enumerate(rows[:15]):
            if 'Số hóa đơn' in [text(v) for v in row]:
                break
        else:
            raise ValueError(f'{sheet}: invoice subheader not found.')
        headers = [text(v) or text(rows[index-1][i]) for i, v in enumerate(row)]
        def col(name):
            if name not in headers:
                raise ValueError(f'{sheet}: missing MISA column {name}')
            return headers.index(name)
        number = col('Số hóa đơn')
        issued = col('Ngày, tháng, năm lập hóa đơn')
        party = col('Mã số thuế người bán' if direction == 'purchases' else 'Mã số thuế người mua')
        net = col('Giá trị HHDV mua vào chưa có thuế' if direction == 'purchases' else 'Doanh thu chưa có thuế GTGT')
        vat = col('Thuế GTGT đủ điều kiện khấu trừ thuế' if direction == 'purchases' else 'Thuế GTGT')
        imported = col('Là chứng từ nhập khẩu HHDV') if direction == 'purchases' else None
        for source_row, row in enumerate(rows[index+1:], index+2):
            if not text(row[number]) or text(row[number]).startswith('##'):
                continue
            notes = text(row[8]) if direction == 'sales' else text(row[12]) if len(row)>12 else ''
            is_import = text(row[imported]).upper() == 'X' if imported is not None else False
            records.append([direction, identifier(row[party]), identifier(row[number], True), '',
                            row[issued], row[net], row[vat], misa_record_status(notes, is_import, customs_as_new, direction),
                            is_import,
                            str(Path(path).resolve()), sheet, source_row])
            extra.append({'notes': notes,
                          'import_marker_k': text(row[10]) if direction == 'purchases' else '',
                          'import_detail_l': text(row[11]) if direction == 'purchases' else '',
                          'status_basis': 'Customs declaration classification policy' if is_import else
                          ('MISA Column I status note' if direction == 'sales' else 'Purchase special-case note; status not inferred') if notes
                          else 'Blank status note: New Invoice'})
    if declared_month != month:
        raise ValueError(f'MISA declaration period is {declared_month}, requested {month}.')
    if found_directions != {'purchases', 'sales'}:
        raise ValueError('Both MISA purchase and sales schedules are required.')
    frame = finish(records)
    # Require both MISA components; missing MISA amounts are not silently zeroed.
    frame['total_amount'] = frame[['net_amount', 'vat_amount']].sum(axis=1, min_count=2)
    for column in ('notes', 'import_marker_k', 'import_detail_l', 'status_basis'):
        frame[column] = [item[column] for item in extra]
    return frame


def find_export(folder, filename):
    """Accept both first-letter-uppercase and legacy lowercase names, also on Linux."""
    folder = Path(folder)
    matches = [p for p in folder.glob('*.xlsx') if p.name.casefold() == filename.casefold()]
    if len(matches) > 1:
        raise ValueError(f'Ambiguous filenames for {filename}: keep only one copy.')
    return matches[0] if matches else folder / (filename[:1].upper() + filename[1:])


def discover_gdt(project_dir, month, tax_id):
    """Choose only a complete exporter run, or the portable manual-input folder."""
    import json
    parent = Path(project_dir) / 'exports' / tax_id / month
    for report in sorted(parent.glob('*/run_report.json'), reverse=True):
        metadata = json.loads(report.read_text(encoding='utf-8'))
        if metadata.get('status') == 'complete':
            return report.parent
    return Path(project_dir) / 'data' / 'gdt' / month


def load_gdt_folder(folder, month, tax_id):
    """Load three purchase files plus combined sales or a legacy electronic file."""
    folder = Path(folder)
    frames, diagnostics = [], []
    for category in ('purchases_coded', 'purchases_no_code', 'purchases_cash_register', 'sales'):
        direction = 'sales' if category == 'sales' else 'purchases'
        combined = find_export(folder, f'sales_all_{month}.xlsx')
        path = combined if category == 'sales' and combined.exists() else find_export(folder, f'{"sales_electronic" if category == "sales" else category}_{month}.xlsx')
        try:
            frame = read_gdt(path, direction, tax_id, month,
                             ['Electronic', 'Cash register'] if path == combined else None)
            frame['gdt_category'] = category
            frames.append(frame)
            diagnostics.append({'category': category, 'file': str(path), 'status': 'loaded', 'rows': len(frame),
                                'detail': 'Both sales types' if path == combined else 'Electronic sales only' if category == 'sales' else ''})
        except (ValueError, OSError) as exc:
            diagnostics.append({'category': category, 'file': str(path), 'status': 'blocked', 'rows': None, 'detail': str(exc)})
    loaded = pd.concat(frames, ignore_index=True) if frames else finish([])
    if 'gdt_category' not in loaded:
        loaded['gdt_category'] = pd.Series(dtype='str')
    return loaded, pd.DataFrame(diagnostics)


def status_summary(frame):
    status = frame['status'].fillna('').str.casefold()
    masks = {'new': status.eq('hóa đơn mới'), 'adjustment_related': status.str.contains('điều chỉnh'),
             'replacement_related': status.str.contains('thay thế'), 'cancelled': status.str.contains('hủy|huỷ|xóa bỏ')}
    flags = pd.DataFrame(masks)
    detail = frame.copy()
    detail['status_category'] = flags.apply(lambda row: next((k for k, v in row.items() if v), 'other'), axis=1)
    detail['overlapping_labels'] = flags.sum(axis=1).gt(1)
    summary = detail['status_category'].value_counts().reindex([*masks, 'other'], fill_value=0).rename_axis('category').reset_index(name='count')
    return summary, detail


def comparison_identifier(value):
    """Ignore numeric leading-zero padding in a comparison key, never in source data."""
    original = identifier(value)
    return (original.lstrip('0') or '0') if original.isdigit() else original


def reconcile(gdt, misa, month, direction, tolerance=1, review_statuses=False):
    if tolerance < 0:
        raise ValueError('Tolerance must be nonnegative.')
    g = gdt.loc[gdt.direction.eq(direction)].copy()
    m = misa.loc[misa.direction.eq(direction)].copy()
    for data in (g, m):
        if data.empty and 'total_amount' not in data:
            data['total_amount'] = pd.Series(dtype=float)
    if 'total_amount' not in g or 'total_amount' not in m:
        raise ValueError('Reload GDT/MISA records: total payment is required for reconciliation.')
    # Both original identifiers remain available independently after the merge.
    keys = ['_key_party_tax_id', '_key_invoice_number', '_key_issue_date']
    for frame in (g, m):
        frame['_key_party_tax_id'] = frame.party_tax_id.map(comparison_identifier)
        frame['_key_invoice_number'] = frame.invoice_number.map(comparison_identifier)
        frame['_key_issue_date'] = frame.issue_date.dt.normalize()
    output = []
    def excluded(row, reason, source):
        record = {f'{source}_{k}': v for k, v in row.items()}
        record.update(result=reason, direction=direction, confidence='review_required')
        output.append(record)
    for source, frame in [('gdt', g), ('misa', m)]:
        frame['eligible'] = True
        for index, row in frame.iterrows():
            reason = None
            if pd.isna(row.issue_date) or not row.party_tax_id or not row.invoice_number:
                reason = 'missing_matching_fields'
            elif review_statuses and source == 'gdt' and re.search('hủy|huỷ|xóa bỏ|điều chỉnh|thay thế', row.status, re.I):
                reason = 'invoice_status_requires_review'
            if reason:
                excluded(row, reason, source)
                frame.at[index, 'eligible'] = False
        duplicates = frame.loc[frame.eligible].duplicated(keys, keep=False)
        for index in duplicates.index[duplicates]:
            excluded(frame.loc[index], 'ambiguous_matching_key', source)
            frame.at[index, 'eligible'] = False
    # Exclude keys ambiguous on either side to prevent a misleading one-sided result.
    ambiguous = set(tuple(row[k] for k in keys) for source, frame in [('gdt', g), ('misa', m)]
                    for _, row in frame.loc[~frame.eligible].iterrows())
    for source, frame in [('gdt', g), ('misa', m)]:
        for index, row in frame.loc[frame.eligible].iterrows():
            if tuple(row[k] for k in keys) in ambiguous:
                excluded(row, 'counterpart_requires_review', source)
                frame.at[index, 'eligible'] = False
    merged = g.loc[g.eligible].merge(m.loc[m.eligible], on=keys, how='outer', suffixes=('_gdt', '_misa'), indicator=True, validate='one_to_one')
    for _, row in merged.iterrows():
        record = {**row.to_dict(), 'direction': direction, 'confidence': 'matched_without_misa_series'}
        if row['_merge'] != 'both':
            record['result'] = 'gdt_only' if row['_merge'] == 'left_only' else 'misa_only'
        else:
            # Treat blank GDT amounts as zero for comparison; keep source blanks intact.
            gdt_total = 0 if pd.isna(row.total_amount_gdt) else row.total_amount_gdt
            gdt_vat = 0 if pd.isna(row.vat_amount_gdt) else row.vat_amount_gdt
            total_delta = row.total_amount_misa - gdt_total
            vat_delta = row.vat_amount_misa - gdt_vat
            net_delta = row.net_amount_misa - row.net_amount_gdt
            record.update(net_difference=net_delta, total_difference=total_delta, vat_difference=vat_delta,
                          gdt_blank_total_as_zero=pd.isna(row.total_amount_gdt),
                          gdt_blank_vat_as_zero=pd.isna(row.vat_amount_gdt))
            if pd.isna(row.net_amount_gdt):
                # Project review category; the missing net amount stays missing.
                record['result'] = 'commercial_invoice_requires_review'
            elif pd.isna(net_delta) or pd.isna(total_delta) or pd.isna(vat_delta):
                record['result'] = 'missing_amount_requires_review'
            elif abs(net_delta) > tolerance or abs(total_delta) > tolerance or abs(vat_delta) > tolerance:
                record['result'] = 'amount_difference_requires_review'
            else:
                record['result'] = 'matched_amounts_agree'
        output.append(record)
    return pd.DataFrame(output)


def display_records(frame):
    """A readable view; source DataFrame column names stay stable for processing."""
    direction = frame.direction.iloc[0] if not frame.empty and frame.direction.nunique() == 1 else None
    labels = {'direction': 'Invoice Type', 'party_tax_id': tax_code_label(direction),
              'invoice_number': 'Invoice Number', 'series': 'Invoice Series', 'issue_date': 'Issue Date',
              'net_amount': 'Amount Excluding VAT', 'vat_amount': 'VAT Amount', 'total_amount': 'Invoice Total', 'status': 'Source Status',
              'notes': 'MISA Notes', 'import_document': 'Import Document',
              'import_marker_k': 'Import Marker (Column K)', 'import_detail_l': 'Import Detail (Column L)',
              'source_file': 'File Path', 'source_sheet': 'Source Sheet', 'source_row': 'Source Row'}
    columns = [c for c in labels if c in frame.columns]
    result = frame[columns].copy()
    if 'direction' in result:
        result['direction'] = result.direction.map({'purchases': 'Input VAT Invoices', 'sales': 'Output VAT Invoices'})
    if 'issue_date' in result:
        result['issue_date'] = result.issue_date.dt.strftime('%d/%m/%Y')
    for column in ('status', 'notes'):
        if column in result:result[column] = result[column].map(english_invoice_status)
    return result.rename(columns=labels)


def source_status_overview(gdt, misa):
    rows = []
    for source, frame in [('GDT', gdt), ('MISA', misa)]:
        for direction, name in [('purchases', 'Input VAT Invoices'), ('sales', 'Output VAT Invoices')]:
            records = frame.loc[frame.direction.eq(direction)]
            counts, _ = status_summary(records)
            values = counts.set_index('category')['count']
            rows.append({'Source': source, 'Invoice Type': name, 'Total Invoices': len(records),
                         **{label: int(values.get(key, 0)) for key, label in STATUS_LABELS.items()}})
    return pd.DataFrame(rows)


def source_review_views(result):
    """Represent every non-agreeing source row once, separately for GDT and MISA."""
    views = {'GDT': [], 'MISA': []}
    reasons = {'gdt_only': 'GDT Only', 'misa_only': 'MISA Only',
               'commercial_invoice_requires_review': 'Commercial Invoices',
               'missing_amount_requires_review': 'Missing Amount',
               'amount_difference_requires_review': 'Amount Difference',
               'invoice_status_requires_review': 'Invoice Status Review',
               'counterpart_requires_review': 'Counterpart Requires Review',
               'import_document_outside_domestic_gdt_scope': 'Customs / Import Document',
               'issue_date_outside_gdt_month': 'Issue Date Outside GDT Month',
               'missing_matching_fields': 'Missing Matching Fields',
               'ambiguous_matching_key': 'Ambiguous Matching Key'}
    for _, row in result.iterrows():
        if row['result'] == 'matched_amounts_agree':
            continue
        def value(source, field):
            for key in (f'{field}_{source}', f'{source}_{field}', field):
                if key in row and not pd.isna(row[key]):
                    return row[key]
            return None
        for source in ('gdt', 'misa'):
            if not value(source, 'source_file'):
                continue
            issued = value(source, 'issue_date')
            reason = reasons.get(row['result'], row['result'])
            document_number = identifier(value(source, 'invoice_number'))
            if source == 'misa' and row.get('direction') == 'purchases' and re.fullmatch(r'\d{12}', document_number):
                reason = 'Custom Declaration'
            views[source.upper()].append({
                'Review Reason': reason,
                tax_code_label(row.get('direction')): value(source, 'party_tax_id'),
                'Invoice Number': value(source, 'invoice_number'),
                'Invoice Series': value(source, 'series'),
                'Issue Date': issued.strftime('%d/%m/%Y') if pd.notna(issued) else None,
                'GDT Amount Excluding VAT': value('gdt', 'net_amount'), 'MISA Amount Excluding VAT': value('misa', 'net_amount'),
                'Net Amount Difference (MISA − GDT)': row.get('net_difference'),
                'GDT Invoice Total': value('gdt', 'total_amount'), 'MISA Invoice Total': value('misa', 'total_amount'),
                'Invoice Total Difference (MISA − GDT)': row.get('total_difference'),
                'GDT VAT Amount': value('gdt', 'vat_amount'), 'MISA VAT Amount': value('misa', 'vat_amount'),
                'VAT Difference (MISA − GDT)': row.get('vat_difference'),
                'Invoice Status / Notes': english_invoice_status(value(source, 'status')),
                'Source Sheet': value(source, 'source_sheet'), 'Source Row': value(source, 'source_row')})
    return {name: pd.DataFrame(rows) for name, rows in views.items()}


def matching_overview(reconciliation, gdt, misa):
    rows = []
    for direction, result in reconciliation.items():
        agrees = int(result['result'].eq('matched_amounts_agree').sum()) if not result.empty else 0
        key_pairs = int(result['result'].isin(PAIRED_RESULTS).sum()) if not result.empty else 0
        reviews = source_review_views(result)
        for source, data in [('GDT', gdt), ('MISA', misa)]:
            total = int(data.direction.eq(direction).sum())
            review = len(reviews[source])
            if agrees + review != total:
                raise ValueError(f'{source}/{direction}: reconciliation source counts do not balance.')
            rows.append({'Source': source, 'Invoice Type': 'Input VAT Invoices' if direction == 'purchases' else 'Output VAT Invoices',
                         'Total Invoices': total, 'Invoices Matched by Key': key_pairs,
                         'Matched Invoices': agrees, 'Amount Differences / Missing Values': key_pairs - agrees,
                         'Unmatched / Duplicate-Key Records': total - key_pairs,
                         'Records Requiring Review': review})
    return pd.DataFrame(rows)


STATUS_PALETTES = {
    'modern': dict(zip(STATUS_LABELS.values(), ['#326B77', '#D9822B', '#6D3507', '#1E293B', '#CBD5E1'])),
    'current': dict(zip(STATUS_LABELS.values(), ['#237C73', '#D99332', '#527DA8', '#B85E65', '#A6ABB2'])),
    'clear': dict(zip(STATUS_LABELS.values(), ['#4477AA', '#E69F00', '#7B4AB5', '#CC3311', '#999999'])),
    'soft': dict(zip(STATUS_LABELS.values(), ['#287D78', '#BE8D35', '#8064A2', '#B85E65', '#939BA5'])),
}


def plot_comparison(overview, month, matching=False, palette="modern"):
    """Matplotlib source-level counts; review records are not labelled tax errors."""
    import matplotlib.pyplot as plt
    import numpy as np
    from matplotlib.colors import to_rgb
    columns = ['Matched Invoices', 'Amount Differences / Missing Values', 'Unmatched / Duplicate-Key Records'] if matching else list(STATUS_LABELS.values())
    if palette not in STATUS_PALETTES:raise ValueError(f'Unknown invoice-status palette: {palette}')
    colors = ['#237C73', '#D99332', '#527DA8'] if matching else [STATUS_PALETTES[palette][column] for column in columns]
    opacity = .95 if palette == 'modern' else 1
    if matching and palette == 'modern':
        colors = ['#A0AEC0', '#EA580C', '#991B1B']
    fig, axes = plt.subplots(1, 2, figsize=(11.5, 5.0), facecolor='#FFFFFF')
    for axis, name in zip(axes, ['Input VAT Invoices', 'Output VAT Invoices']):
        axis.set_facecolor('#FFFFFF')
        data = overview.loc[overview['Invoice Type'].eq(name)].set_index('Source').reindex(['GDT', 'MISA'])
        bottom = np.zeros(2)
        maximum = max(float(data[columns].sum(axis=1).max()), 1)
        small_segments = []
        for column, color in zip(columns, colors):
            values = data[column].to_numpy(dtype=float)
            legend_label = 'Matching Amounts' if matching and column == 'Matched Invoices' else column
            if matching and column == 'Amount Differences / Missing Values':legend_label='Amount Diff / Missing Values'
            bars = axis.bar(['GDT', 'MISA'], values, bottom=bottom, label=legend_label, color=color, alpha=opacity, edgecolor='none', width=.55)
            r, g, b = to_rgb(color)
            lum = 0.2126*r + 0.7152*g + 0.0722*b
            inner_text_c = '#0F172A' if lum > 0.65 else '#FFFFFF'
            for x, (bar, value) in enumerate(zip(bars, values)):
                if not value:continue
                center = bar.get_y() + bar.get_height() / 2
                if value < maximum * .04:
                    short = {'Amount Differences / Missing Values': 'Amount Diff', 'Unmatched / Duplicate-Key Records': 'Unmatched', 'Matched Invoices': 'Matched'}.get(column, column.replace(' Invoices', ''))
                    label_color = '#64748B' if column in ['Other', 'Matched Invoices'] or color in ['#CBD5E1', '#A0AEC0'] else ('#A05206' if color == '#D9822B' else color)
                    small_segments.append((x, center, f'{short}: {int(value)}', color, label_color))
                else:
                    axis.text(x, center, str(int(value)), ha='center', va='center', fontsize=9.5,
                              color=inner_text_c, fontweight='bold',
                              bbox={'facecolor': color, 'alpha': 0.15, 'edgecolor': 'none', 'pad': 2, 'boxstyle': 'round,pad=0.25'})
            bottom += values
        for x, total in enumerate(bottom):
            axis.text(x, total + max(bottom.max() * .028, .4), f'Total: {int(total)}', ha='center', fontsize=10, fontweight='bold', color='#0F172A',zorder=6,bbox={'facecolor':'white','edgecolor':'none','pad':.8})
        slots = {0: 0, 1: 0}
        for x, center, label, color, label_color in small_segments:
            level = slots[x];slots[x] += 1
            callout_x = 0.28 if x == 0 else 1.52
            axis.annotate(label, xy=(x + .26, center), xytext=(callout_x, maximum * (1.13 + .11 * level)),
                          ha='left', fontsize=8.5, color=label_color, fontweight='bold',
                          zorder=4,arrowprops={'arrowstyle': '-', 'color': color, 'linewidth': 1.1, 'relpos': (0, 0),'shrinkA':3,'shrinkB':2},
                          bbox={'facecolor': '#FFFFFF', 'edgecolor': color, 'pad': 2.2, 'boxstyle': 'round,pad=0.25', 'linewidth': 1.0})
        if small_segments:
            axis.set_xlim(-.5, 2.55)
        headroom = max(1.22, 1.25 + .11 * (max(slots.values()) - 1)) if small_segments else 1.22
        axis.set_ylim(0, max(maximum * headroom, 1))
        axis.set_title(name, fontsize=12, fontweight='bold', color='#0F172A', pad=12)
        axis.set_ylabel('Source Records', fontsize=10, color='#475569')
        axis.spines[['top', 'right']].set_visible(False)
        axis.spines[['left', 'bottom']].set_color('#CBD5E1')
        axis.tick_params(colors='#475569', labelsize=9.5)
        axis.yaxis.set_major_locator(plt.MaxNLocator(integer=True))
    handles, labels = axes[0].get_legend_handles_labels()
    legend = fig.legend(handles, labels, loc='upper center', bbox_to_anchor=(.5, .93), ncol=2 if matching else 3, frameon=False, fontsize=9.5)
    for label, color in zip(legend.get_texts(), colors):
        legend_color = '#64748B' if color.upper() in ['#CBD5E1', '#EBEBD3', '#A0AEC0'] else ('#A05206' if color == '#D9822B' else color)
        label.set_color(legend_color)
        label.set_fontweight('bold')
    title_text = f'{month} | ' + ('Invoice Reconciliation Results' if matching else 'Invoice Records by Source and Status')
    fig.suptitle(title_text, fontsize=13.5, fontweight='bold', color='#0F172A', y=.99)
    fig.text(.5, .015, 'Matching key: invoice number, issue date and seller’s/purchaser’s tax code. Each matched pair is counted once per source.' if matching else 'Status counts include related original invoices; customs declarations remain separately identified.', ha='center', fontsize=8.5, color='#475569')
    fig.subplots_adjust(top=.72, bottom=.15, wspace=.28)
    # Round the outer stack, keeping internal status boundaries flat. The radius
    # is 10% of bar width; convert the vertical radius to preserve circular corners.
    from matplotlib.path import Path as PlotPath
    fig.canvas.draw()
    for axis in axes:
        for x in (0, 1):
            segments = [bar for bar in axis.patches
                        if abs(bar.get_x() + bar.get_width() / 2 - x) < 1e-8]
            if not segments: continue
            total = max(bar.get_y() + bar.get_height() for bar in segments)
            if total <= 0: continue
            width = segments[0].get_width()
            left, right = x - width / 2, x + width / 2
            x_pixels = axis.bbox.width / (axis.get_xlim()[1] - axis.get_xlim()[0])
            y_pixels = axis.bbox.height / (axis.get_ylim()[1] - axis.get_ylim()[0])
            radius_x = min(width * .10, total * y_pixels / (2 * x_pixels))
            radius_y = radius_x * x_pixels / y_pixels
            vertices = [(left+radius_x,0),(right-radius_x,0),
                        (right,0),(right,radius_y),(right,total-radius_y),
                        (right,total),(right-radius_x,total),(left+radius_x,total),
                        (left,total),(left,total-radius_y),(left,radius_y),
                        (left,0),(left+radius_x,0),(left+radius_x,0)]
            codes = [PlotPath.MOVETO,PlotPath.LINETO,
                     PlotPath.CURVE3,PlotPath.CURVE3,PlotPath.LINETO,
                     PlotPath.CURVE3,PlotPath.CURVE3,PlotPath.LINETO,
                     PlotPath.CURVE3,PlotPath.CURVE3,PlotPath.LINETO,
                     PlotPath.CURVE3,PlotPath.CURVE3,PlotPath.CLOSEPOLY]
            outline = PlotPath(vertices, codes)
            for bar in segments: bar.set_clip_path(outline, axis.transData)
    return fig


def save_chart(figure, path):
    """Write a complete PNG atomically, avoiding partially written images in synced folders."""
    from io import BytesIO
    buffer = BytesIO()
    figure.savefig(buffer, format='png', dpi=160, bbox_inches='tight')
    path = Path(path)
    temporary = path.with_name(path.stem + '.pending.png')
    temporary.write_bytes(buffer.getvalue())
    temporary.replace(path)


def reconciliation_export_tables(result):
    """Partition every comparison/review row into the four requested worksheets.

    Mismatch includes scope, status and ambiguous-key reviews, with explicit reasons;
    unpaired invoices go in the two source-only worksheets, including records without usable keys.
    """
    sheet_results = {
        'Matching Reconciliations': {'matched_amounts_agree'},
        'Mismatch Reconciliations': set(result.get('result', [])) - {'matched_amounts_agree', 'gdt_only', 'misa_only', 'missing_matching_fields'},
        'Invoices Only From GDT': {'gdt_only'},
        'Invoices Only From MISA': {'misa_only'},
    }
    labels = {'party_tax_id': tax_code_label(result.direction.iloc[0] if not result.empty else None), 'invoice_number': 'Invoice Number',
              'series': 'Invoice Series', 'issue_date': 'Issue Date', 'net_amount': 'Amount Excluding VAT',
              'vat_amount': 'VAT Amount', 'total_amount': 'Invoice Total', 'status': 'Source Status', 'import_document': 'Import Document',
              'source_file': 'File Path', 'source_sheet': 'Source Sheet', 'source_row': 'Source Row'}
    columns = ['Result'] + [f'{source.upper()} {label}' for source in ('gdt', 'misa') for label in labels.values()] + ['Net Amount Difference', 'Invoice Total Difference', 'VAT Difference']
    tables = {}
    for name, statuses in sheet_results.items():
        records = []
        for _, row in result.iterrows():
            if row['result'] not in statuses:
                if not (row['result'] == 'missing_matching_fields' and ((name == 'Invoices Only From GDT' and pd.notna(row.get('gdt_source_file'))) or (name == 'Invoices Only From MISA' and pd.notna(row.get('misa_source_file'))))):
                    continue
            record = {'Result': 'Commercial Invoices' if row['result'] == 'commercial_invoice_requires_review' else row['result'].replace('_', ' ').title()}
            for source in ('gdt', 'misa'):
                # Shared matching fields must not invent a record on an absent side.
                present = any(pd.notna(row.get(key)) for key in (f'source_file_{source}', f'{source}_source_file'))
                for field, label in labels.items():
                    value = None
                    if present:
                        for key in (f'{field}_{source}', f'{source}_{field}', field):
                            if key in row and pd.notna(row[key]):
                                value = row[key]
                                break
                    record[f'{source.upper()} {label}'] = value
            record['Net Amount Difference'] = row.get('net_difference')
            record['Invoice Total Difference'] = row.get('total_difference')
            record['VAT Difference'] = row.get('vat_difference')
            records.append(record)
        tables[name] = pd.DataFrame(records, columns=columns)
    if sum(map(len, tables.values())) != len(result):
        raise ValueError('Export partition does not account for every result.')
    return tables


def export_reconciliation_workbook(result, path, source_layouts=True):
    """Four sheets: concise matches, vertically stacked mismatches, original-only tables."""
    from openpyxl import Workbook
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
    book = Workbook()
    book.remove(book.active)
    tables = original_header_export_tables(result) if source_layouts else reconciliation_export_tables(result)
    def write_table(sheet, frame, start=1):
        visible = [re.sub(r'__blank_column_\d+', '', header).strip() or None for header in frame.columns]
        for column, header in enumerate(visible, 1):
            cell = sheet.cell(start, column, header)
            cell.font = Font(bold=True, color='254C5A')
            cell.fill = PatternFill('solid', fgColor='D7E7EA')
            cell.alignment = Alignment(wrap_text=True, vertical='top')
        sheet.row_dimensions[start].height = 42
        for offset, record in enumerate(frame.itertuples(index=False, name=None), start + 1):
            for column, (header, value) in enumerate(zip(frame.columns, record), 1):
                if pd.isna(value):continue
                cell = sheet.cell(offset, column, value)
                is_id = any(term in header for term in ('Tax ID', 'Tax Code', 'Invoice Number', 'Invoice Series', 'MST', 'Mã số thuế', 'Số hóa đơn', 'Ký hiệu hóa đơn'))
                is_amount = any(term in header for term in ('Amount', 'Payment', 'Difference', 'Tổng tiền', 'Giá trị HHDV', 'Doanh thu', 'Thuế GTGT'))
                if is_id:
                    cell.value = identifier(value);cell.data_type = 's';cell.number_format = '@'
                elif isinstance(value, str):cell.data_type = 's'
                elif any(term in header for term in ('Issue Date', 'Ngày', 'ngày')) and not isinstance(value, (int, float)):
                    cell.number_format = 'dd/mm/yyyy'
                elif is_amount:cell.number_format = '#,##0;[Red](#,##0);"-"'
                width = 45 if 'File Path' in header else 28
                wrap = isinstance(value, str) and len(value) > width
                cell.alignment = Alignment(horizontal='right' if is_amount else 'left', vertical='top', wrap_text=wrap)
                if wrap:
                    import math
                    height = min(90, 15 * math.ceil(len(value) / width))
                    sheet.row_dimensions[offset].height = max(sheet.row_dimensions[offset].height or 15, height)
        for column, header in enumerate(frame.columns, 1):
            width = 45 if 'File Path' in header else 28
            sheet.column_dimensions[get_column_letter(column)].width = max(sheet.column_dimensions[get_column_letter(column)].width or 0, width)
        return start + len(frame) + 1
    for name, frame in tables.items():
        sheet = book.create_sheet(name)
        if name == 'Mismatch Reconciliations' and source_layouts:
            blocks = frame.attrs.get('vertical_blocks', [])
            if not blocks:sheet.cell(1, 1, 'No mismatches.')
            line = 1
            for block in blocks:
                width = max(len(source['table'].columns) for source in block['sources'])
                for column in range(1, width + 1):
                    cell = sheet.cell(line, column)
                    cell.fill = PatternFill('solid', fgColor='315969')
                    cell.font = Font(bold=True, color='FFFFFF')
                sheet.cell(line, 1, f'Review {block["number"]}: {block["reason"]}')
                sheet.row_dimensions[line].height = 22
                line += 1
                for source in block['sources']:
                    sheet.cell(line, 1, source['source']).font = Font(bold=True)
                    sheet.cell(line + 1, 1, f'{source["file"]} | Sheet: {source["sheet"]} | Source Row: {source["row"]}').data_type = 's'
                    line = write_table(sheet, source['table'], line + 2) + 1
                def amount(value):return 'Unavailable' if pd.isna(value) else f'{value:,.0f}'
                sheet.cell(line, 1, f'Difference (MISA - GDT): Net Amount {amount(block["net_difference"])}; Invoice Total {amount(block["total_difference"])}; VAT {amount(block["vat_difference"])}')
                line += 3
        else:
            write_table(sheet, frame)
            sheet.freeze_panes = 'A2'
            sheet.auto_filter.ref = sheet.dimensions
    path = Path(path);path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_name(path.stem + '.pending.xlsx')
    book.save(temporary);book.close();temporary.replace(path)
    return path


def original_header_export_tables(result):
    """Use complete original source columns; prefixes distinguish the paired sides.

    The standard GDT order is retained, with cash-register-only columns appended.
    MISA's two header levels are flattened using its invoice subheader and parent
    labels, in the original column order. No source invoice value is normalised.
    """
    files, layouts = {}, {}
    def reference(row, source, field):
        for key in (f'{field}_{source}', f'{source}_{field}'):
            value = row.get(key)
            if value is not None and pd.notna(value):
                return value
        return None
    refs = {source: [] for source in ('gdt', 'misa')}
    for _, row in result.iterrows():
        for source in refs:
            path, sheet = reference(row, source, 'source_file'), reference(row, source, 'source_sheet')
            if path is not None and (str(path), sheet) not in refs[source]:
                refs[source].append((str(path), sheet))
    # Stable layout: coded / uncoded exports before cash-register-specific columns.
    refs['gdt'].sort(key=lambda item: ('cash_register' in Path(item[0]).name.lower(), item[0], item[1]))
    headers = {'gdt': [], 'misa': []}
    for source, references in refs.items():
        for path, sheet in references:
            if path not in files:
                files[path] = read_book(path)
            rows = files[path][sheet]
            for index, row in enumerate(rows[:30]):
                if 'Số hóa đơn' in [text(value) for value in row]:
                    break
            else:
                raise ValueError(f'{path}/{sheet}: original invoice headers not found.')
            raw_headers = list(rows[index])
            if source == 'misa' and index:
                raw_headers = [value if text(value) else rows[index-1][i] for i, value in enumerate(raw_headers)]
            nonblank = [i for i, value in enumerate(raw_headers) if text(value)]
            width = max(nonblank) + 1
            # MISA contains merged-header spacer columns; retain their positions
            # and values without inventing visible field names.
            names = [str(value) if text(value) else f'__blank_column_{i+1}' for i, value in enumerate(raw_headers[:width])]
            if len(names) != len(set(names)):
                raise ValueError(f'{path}/{sheet}: duplicate original headers need explicit mapping.')
            layouts[(path, sheet)] = (names, rows)
            for name in names:
                if name not in headers[source]:headers[source].append(name)
    def source_record(row, source):
        path = reference(row, source, 'source_file')
        if path is None:return {}
        sheet, number = reference(row, source, 'source_sheet'), reference(row, source, 'source_row')
        names, rows = layouts[(str(path), sheet)]
        record = dict(zip(names, rows[int(number)-1]))
        if source == 'gdt' and row.get('direction') == 'purchases':
            category = reference(row, source, 'gdt_category')
            if not category:
                filename = Path(path).name.lower()
                category = 'purchases_cash_register' if 'cash_register' in filename else 'purchases_no_code' if 'no_code' in filename else 'purchases_coded'
            record['GDT Invoice Type'] = {'purchases_coded': 'Coded', 'purchases_no_code': 'No Code', 'purchases_cash_register': 'Cash Register'}.get(category, category)
        return record
    is_input = not result.empty and result.direction.eq('purchases').all()
    gdt_headers = headers['gdt'] + (['GDT Invoice Type'] if is_input else [])
    paired_columns = (['Result'] + [f'GDT {h}' if h != 'GDT Invoice Type' else h for h in gdt_headers]
                      + ['GDT File Path', 'GDT Source Sheet', 'GDT Source Row']
                      + [f'MISA {h}' for h in headers['misa']]
                      + ['MISA File Path', 'MISA Source Sheet', 'MISA Source Row', 'Net Amount Difference', 'Invoice Total Difference', 'VAT Difference'])
    buckets = {'Matching Reconciliations': [], 'Mismatch Reconciliations': [], 'Invoices Only From GDT': [], 'Invoices Only From MISA': []}
    for _, row in result.iterrows():
        if row['result'] in ('gdt_only', 'misa_only', 'missing_matching_fields'):
            source = ('gdt' if reference(row, 'gdt', 'source_file') is not None else 'misa')
            buckets['Invoices Only From ' + source.upper()].append(source_record(row, source))
            continue
        record = {'Result': 'Commercial Invoices' if row['result'] == 'commercial_invoice_requires_review' else row['result'].replace('_', ' ').title()}
        for source in ('gdt', 'misa'):
            for label, value in source_record(row, source).items():
                record[label if label == 'GDT Invoice Type' else f'{source.upper()} {label}'] = value
            for field, label in [('source_file', 'File Path'), ('source_sheet', 'Source Sheet'), ('source_row', 'Source Row')]:
                record[f'{source.upper()} {label}'] = reference(row, source, field)
        record['Net Amount Difference'] = row.get('net_difference')
        record['Invoice Total Difference'], record['VAT Difference'] = row.get('total_difference'), row.get('vat_difference')
        bucket = 'Matching Reconciliations' if row['result'] == 'matched_amounts_agree' else 'Mismatch Reconciliations'
        buckets[bucket].append(record)
    columns = {'Matching Reconciliations': paired_columns, 'Mismatch Reconciliations': paired_columns,
               'Invoices Only From GDT': gdt_headers, 'Invoices Only From MISA': headers['misa']}
    tables = {name: pd.DataFrame(records, columns=columns[name]) for name, records in buckets.items()}
    tables['Mismatch Reconciliations'].attrs['vertical_blocks'] = mismatch_blocks(tables['Mismatch Reconciliations'], gdt_headers, headers['misa'])
    tables['Matching Reconciliations'] = consolidate_matching_table(tables['Matching Reconciliations'], is_input=is_input)
    return tables


def consolidate_matching_table(paired, is_input=None):
    """One column per equivalent field; retain source labels when values differ."""
    if is_input is None:
        is_input = 'MISA Giá trị HHDV mua vào chưa có thuế' in paired
    role = 'bán' if is_input else 'mua'
    gdt_party = next((c for c in paired if c.startswith(f'GDT MST người {role}')), None)
    gdt_name = next((c for c in paired if c.startswith(f'GDT Tên người {role}')), None)
    misa_net = 'MISA ' + ('Giá trị HHDV mua vào chưa có thuế' if is_input else 'Doanh thu chưa có thuế GTGT')
    misa_vat = 'MISA ' + ('Thuế GTGT đủ điều kiện khấu trừ thuế' if is_input else 'Thuế GTGT')
    fields = [
        ('Invoice Number', 'GDT Số hóa đơn', 'MISA Số hóa đơn', 'identifier'),
        ('Issue Date', 'GDT Ngày lập', 'MISA Ngày, tháng, năm lập hóa đơn', 'date'),
        (f"{'Seller' if is_input else 'Purchaser'}'s Tax Code", gdt_party, f'MISA Mã số thuế người {role}', 'identifier'),
        (f'{"Seller" if is_input else "Purchaser"} Name', gdt_name, f'MISA Tên người {role}', 'text'),
        ('Amount Excluding VAT', 'GDT Tổng tiền chưa thuế', misa_net, 'amount'),
        ('VAT Amount', 'GDT Tổng tiền thuế', misa_vat, 'vat'),
        ('Invoice Total', 'GDT Tổng tiền thanh toán', None, 'total'),
    ]
    used = {column for _, g, m, _ in fields for column in (g, m) if column}
    used.update({'Result', 'Net Amount Difference', 'Invoice Total Difference', 'VAT Difference', 'GDT STT', 'MISA STT'})
    provenance = [f'{source} {field}' for source in ('GDT', 'MISA') for field in ('File Path', 'Source Sheet', 'Source Row')]
    remaining = [c for c in paired if c not in used and c not in provenance
                 and ('__blank_column_' not in c or paired[c].notna().any())]
    result = pd.DataFrame(index=paired.index)
    def absent(value):return value is None or pd.isna(value) or (isinstance(value, str) and not value.strip())
    def labelled(value, amount=False):
        if absent(value):return '(blank)'
        return f'{value:,.0f}' if amount and isinstance(value, (int, float)) else str(value)
    def combine(g, m, kind):
        if kind == 'identifier':
            values = [identifier(v) for v in (g, m) if not absent(v)]
            return max(values, key=lambda v: len(v) - len(v.lstrip('0')), default='')
        if kind == 'date':
            return pd.to_datetime(g if not absent(g) else m, dayfirst=True).normalize()
        if kind in ('amount', 'vat', 'total'):
            g = pd.to_numeric(g) if not absent(g) else g
            m = pd.to_numeric(m) if not absent(m) else m
        if kind in ('vat', 'total') and absent(g):g = 0
        if absent(g) and absent(m):return None
        if not absent(g) and not absent(m) and g == m:return g
        # Preserve a meaningful difference rather than silently selecting a source.
        return f'GDT: {labelled(g, kind in ("amount", "vat", "total"))} | MISA: {labelled(m, kind in ("amount", "vat", "total"))}'
    for label, g, m, kind in fields:
        values = []
        for _, row in paired.iterrows():
            gv, mv = row.get(g), row.get(m)
            if kind == 'total':
                net, vat = row.get(misa_net), row.get(misa_vat)
                mv = None if absent(net) or absent(vat) else pd.to_numeric(net) + pd.to_numeric(vat)
            values.append(combine(gv, mv, kind))
        result[label] = pd.Series(values, index=paired.index, dtype=object)
    for column in remaining + provenance:
        if column in paired:result[column] = paired[column]
    result.attrs['identifier_columns'] = ["Invoice Number", "Seller's Tax Code", "Purchaser's Tax Code"]
    return result


def mismatch_blocks(paired, gdt_headers, misa_headers):
    """Stack source records, aligning equivalent fields in the same Excel columns."""
    # Keep GDT's order as the anchor. In the standard export, invoice number is D.
    # A MISA field retains its source label but occupies the equivalent GDT column.
    is_input = any('Mã số thuế người bán' in header for header in misa_headers)
    role = 'bán' if is_input else 'mua'
    equivalents = {'STT': 'STT', 'Số hóa đơn': 'Số hóa đơn',
                   'Ngày, tháng, năm lập hóa đơn': 'Ngày lập',
                   'Giá trị HHDV mua vào chưa có thuế': 'Tổng tiền chưa thuế',
                   'Doanh thu chưa có thuế GTGT': 'Tổng tiền chưa thuế',
                   'Thuế GTGT đủ điều kiện khấu trừ thuế': 'Tổng tiền thuế',
                   'Thuế GTGT': 'Tổng tiền thuế'}
    for label, prefix in [(f'Mã số thuế người {role}', f'MST người {role}'),
                          (f'Tên người {role}', f'Tên người {role}')]:
        target = next((h for h in gdt_headers if h.startswith(prefix)), None)
        if target: equivalents[label] = target
    positions = {h: i for i, h in enumerate(gdt_headers)}
    misa_positions = {}
    next_position = len(gdt_headers)
    for header in misa_headers:
        target = equivalents.get(header)
        if target in positions:
            misa_positions[header] = positions[target]
        elif '__blank_column_' not in header:
            misa_positions[header] = next_position
            next_position += 1
    def aligned(values, source):
        columns = [f'__blank_column_{i+1}' for i in range(next_position)]
        content = [None] * next_position
        selected = positions if source == 'GDT' else misa_positions
        for header, value in values.items():
            if header in selected:
                index = selected[header]
                columns[index], content[index] = header, value
        return pd.DataFrame([content], columns=columns)
    blocks = []
    for number, (_, row) in enumerate(paired.iterrows(), 1):
        sources = []
        for source, headers in [('GDT', gdt_headers), ('MISA', misa_headers)]:
            path = row.get(f'{source} File Path')
            if path is None or pd.isna(path):continue
            values = {h: row.get(h if h == 'GDT Invoice Type' else f'{source} {h}') for h in headers}
            sources.append({'source': source, 'table': aligned(values, source),
                            'file': path, 'sheet': row.get(f'{source} Source Sheet'), 'row': row.get(f'{source} Source Row')})
        blocks.append({'number': number, 'reason': row.get('Result'),
                       'net_difference': row.get('Net Amount Difference'), 'total_difference': row.get('Invoice Total Difference'),
                       'vat_difference': row.get('VAT Difference'), 'sources': sources})
    return blocks


def tax_code_label(direction):
    """Use the invoice party's role in presentation labels."""
    return {"purchases": "Seller's Tax Code", "sales": "Purchaser's Tax Code"}.get(direction, "Seller's / Purchaser's Tax Code")


def review_block_html(block):
    """Render one vertically stacked review with aligned source columns."""
    from html import escape
    width = max(len(source['table'].columns) for source in block['sources'])
    def label(header):
        return escape(re.sub(r'__blank_column_\d+', '', str(header)))
    parts = [f'<div style="overflow-x:auto;margin:16px 0 28px"><table style="border-collapse:collapse;table-layout:fixed;width:{width*170}px">',
             f'<tr><th colspan="{width}" style="background:#315969;color:white;text-align:left;padding:10px">Review {block["number"]}: {escape(str(block["reason"]))}</th></tr>']
    for source in block['sources']:
        parts.append(f'<tr><td colspan="{width}" style="padding:9px 6px;color:#315969;background:#F3F7F8"><b>{source["source"]}</b> · {escape(str(source["file"]))} · Sheet: {escape(str(source["sheet"]))} · Row: {source["row"]}</td></tr>')
        parts.append('<tr>' + ''.join(f'<th style="background:#D7E7EA;color:#254C5A;text-align:left;padding:8px;overflow-wrap:anywhere">{label(h)}</th>' for h in source['table'].columns) + '</tr>')
        values = []
        for header, value in source['table'].iloc[0].items():
            if value is None or pd.isna(value): value = '-'
            elif isinstance(value, (pd.Timestamp, __import__('datetime').datetime)): value = value.strftime('%d/%m/%Y')
            elif any(term in header for term in ('Tổng tiền', 'Giá trị HHDV', 'Doanh thu', 'Thuế GTGT')): value = format_table_value(value)
            values.append(f'<td style="padding:8px;vertical-align:top;overflow-wrap:anywhere;background:white;color:#253C47">{escape(str(value))}</td>')
        parts.append('<tr>' + ''.join(values) + '</tr>')
    parts.append('</table></div>')
    return ''.join(parts)


def english_invoice_status(value):
    """Translate recognised status labels without rewriting original source records."""
    statuses = {
        'Hóa đơn mới': 'New Invoice',
        'Hóa đơn điều chỉnh': 'Adjustment Invoice',
        'Hóa đơn đã bị điều chỉnh': 'Adjusted Invoice',
        'Hóa đơn thay thế': 'Replacement Invoice',
        'Hóa đơn đã bị thay thế': 'Replaced Invoice',
        'Hóa đơn đã hủy': 'Cancelled Invoice',
        'Hóa đơn đã huỷ': 'Cancelled Invoice',
        'Hóa đơn hủy': 'Cancelled Invoice',
        'Hóa đơn bị xóa bỏ': 'Deleted Invoice',
    }
    return statuses.get(text(value), text(value))


def format_table_value(value):
    """Show numeric zeros as dashes; retain identifiers, dates and source values."""
    from numbers import Real
    if value is None or pd.isna(value):return '-'
    if isinstance(value, Real) and not isinstance(value, bool):
        return '-' if value == 0 else f'{value:,.0f}'
    return value
