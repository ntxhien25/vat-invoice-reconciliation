"""Period, export validation and Windows-only encrypted authentication helpers."""
from __future__ import annotations

import calendar
import ctypes
import json
import os
import re
import unicodedata
from dataclasses import dataclass
from datetime import date, datetime, timedelta, timezone
from pathlib import Path

from openpyxl import Workbook, load_workbook

SAIGON = timezone(timedelta(hours=7))


def export_filename(category: str, month: str) -> str:
    """Capitalise the first character only, retaining the established naming convention."""
    return f"{category[:1].upper()}{category[1:]}_{month}.xlsx"


class ExportError(Exception):
    """A concise, safe error that can be included in the run report."""


@dataclass(frozen=True)
class Period:
    start: date
    end: date

    @classmethod
    def month(cls, value: str) -> "Period":
        if not re.fullmatch(r"\d{4}-\d{2}", value):
            raise ExportError("Month must have format YYYY-MM, for example 2026-09.")
        try:
            year, month = map(int, value.split("-"))
            return cls(date(year, month, 1), date(year, month, calendar.monthrange(year, month)[1]))
        except ValueError as exc:
            raise ExportError("Invalid calendar month.") from exc

    @property
    def label(self) -> str:
        return self.start.strftime("%Y-%m")


def previous_month() -> str:
    return (datetime.now(SAIGON).date().replace(day=1) - timedelta(days=1)).strftime("%Y-%m")


def clean(value) -> str:
    return re.sub(r"\s+", " ", unicodedata.normalize("NFC", str(value or ""))).strip()


def validate_tax_id(value: str) -> str:
    value = value.strip()
    if not re.fullmatch(r"\d{10}(?:-\d{3})?|\d{12}", value):
        raise ExportError("Use a 10-digit MST, a branch MST such as 0123456789-001, or a 12-digit identifier.")
    return value


@dataclass
class InvoiceTable:
    headers: list[str]
    rows: list[tuple]
    duplicate_keys: int = 0


def read_invoice_table(path: Path, expected: int, period: Period, tax_id: str,
                       direction: str, result: str = "Tất cả") -> InvoiceTable:
    """Fail closed on empty/corrupt/wrong-account/stale/page-only exports.

    Source numbers, text identifiers, blanks, negative values and statuses are
    retained. No deduplication or VAT eligibility decision is made here.
    """
    if not path.is_file() or path.stat().st_size == 0:
        raise ExportError("The export is missing or empty.")
    try:
        book = load_workbook(path, read_only=True, data_only=False)
    except Exception as exc:
        raise ExportError("The download is not a readable XLSX workbook.") from exc
    required = {"Số hóa đơn", "Ký hiệu hóa đơn", "Ngày lập"}
    tables = []
    try:
        for sheet in book:
            # GDT can incorrectly declare dimension A1 while storing many rows.
            sheet.reset_dimensions()
            iterator = sheet.iter_rows(values_only=True)
            headers = None
            for number, row in enumerate(iterator, 1):
                names = [clean(v) for v in row]
                if required.issubset(names):
                    headers = names
                    break
                if number >= 30:
                    break
            if headers is None:
                continue
            number_idx = headers.index("Số hóa đơn")
            series_idx = headers.index("Ký hiệu hóa đơn")
            date_idx = headers.index("Ngày lập")
            own_prefix = "MST người mua" if direction == "purchases" else "MST người bán"
            own_indices = [i for i, h in enumerate(headers) if h.startswith(own_prefix)]
            if len(own_indices) != 1:
                raise ExportError("Cannot identify the taxpayer column in the exported workbook.")
            result_idx = next((i for i, h in enumerate(headers)
                               if h in ("Kết quả kiểm tra hóa đơn", "Kết quả xử lý")), None)
            if result != "Tất cả" and result_idx is None:
                raise ExportError("Cannot verify the processing-result category in the workbook.")
            rows, keys = [], set()
            duplicates = 0
            for row in iterator:
                row = tuple(row) + (None,) * max(0, len(headers) - len(row))
                if clean(row[number_idx]) == "Số hóa đơn":
                    continue
                if not clean(row[number_idx]):
                    continue  # title/footer rows; not invoices
                if not clean(row[series_idx]):
                    raise ExportError("An invoice is missing its invoice series.")
                own_id = clean(row[own_indices[0]])
                if own_id != tax_id:
                    raise ExportError("The exported taxpayer does not match the requested MST.")
                value = row[date_idx]
                try:
                    if isinstance(value, datetime):
                        issued = value.date()
                    elif isinstance(value, date):
                        issued = value
                    else:
                        issued = datetime.strptime(clean(value), "%d/%m/%Y").date()
                except (ValueError, TypeError) as exc:
                    raise ExportError("An invoice has an unrecognised issue date.") from exc
                if not period.start <= issued <= period.end:
                    raise ExportError("The workbook contains an invoice outside the requested month.")
                accepted_results = {result.casefold()}
                if result == "Cục Thuế đã nhận không mã":
                    # The XLSX still uses the former authority name for this category.
                    accepted_results.add("tổng cục thuế đã nhận không mã")
                if result != "Tất cả" and clean(row[result_idx]).casefold() not in accepted_results:
                    raise ExportError("The workbook contains a different processing-result category.")
                key = tuple(clean(row[i]) for i in (own_indices[0], series_idx, number_idx))
                # Include counterparty and form to avoid conflating supplier invoice numbers.
                key += tuple(clean(row[i]) for i, h in enumerate(headers)
                             if h.startswith("MST người") and i != own_indices[0] or h == "Ký hiệu mẫu số")
                duplicates += int(key in keys)
                keys.add(key)
                rows.append(tuple(row))
            tables.append(InvoiceTable(headers, rows, duplicates))
    finally:
        book.close()
    if len(tables) != 1:
        raise ExportError("Expected exactly one invoice table; the workbook is blank or its layout changed.")
    table = tables[0]
    if len(table.rows) != expected:
        raise ExportError(f"Portal reports {expected} invoices; workbook has {len(table.rows)}. No complete export confirmed.")
    return table


def write_sales_workbook(path: Path, tables: list[tuple[str, InvoiceTable]], metadata: dict) -> None:
    """Combine validated sales sources into separate sheets, with provenance."""
    book = Workbook()
    book.remove(book.active)
    overview = book.create_sheet("Run information")
    for key, value in metadata.items():
        overview.append((str(key), str(value)))
    overview.append(("Important", "Source values retained; no VAT treatment or deduplication applied."))
    for name, table in tables:
        sheet = book.create_sheet(name)
        sheet.append(table.headers)
        for row in table.rows:
            sheet.append(row)
        # Prevent source text beginning '=' from being interpreted as a new formula.
        for row in sheet:
            for cell in row:
                if isinstance(cell.value, str):
                    cell.data_type = "s"
        sheet.freeze_panes = "A2"
        sheet.auto_filter.ref = sheet.dimensions
    temporary = path.with_suffix(".partial.xlsx")
    book.save(temporary)
    book.close()
    temporary.replace(path)


def dpapi(data: bytes, decrypt: bool = False) -> bytes:
    """Windows DPAPI encryption for the current Windows user; never plaintext."""
    if os.name != "nt":
        raise ExportError("Remembered sessions require Windows DPAPI on this version.")

    class Blob(ctypes.Structure):
        _fields_ = [("length", ctypes.c_ulong), ("data", ctypes.POINTER(ctypes.c_ubyte))]

    buffer = ctypes.create_string_buffer(data)
    source = Blob(len(data), ctypes.cast(buffer, ctypes.POINTER(ctypes.c_ubyte)))
    output = Blob()
    crypt = ctypes.WinDLL("crypt32", use_last_error=True)
    kernel = ctypes.WinDLL("kernel32", use_last_error=True)
    function = crypt.CryptUnprotectData if decrypt else crypt.CryptProtectData
    function.argtypes = [ctypes.POINTER(Blob), ctypes.c_void_p, ctypes.c_void_p,
                         ctypes.c_void_p, ctypes.c_void_p, ctypes.c_ulong, ctypes.POINTER(Blob)]
    function.restype = ctypes.c_int
    kernel.LocalFree.argtypes = [ctypes.c_void_p]
    kernel.LocalFree.restype = ctypes.c_void_p
    if not function(ctypes.byref(source), None, None, None, None, 1, ctypes.byref(output)):
        raise ExportError("Windows could not protect/read the authentication cache.")
    try:
        return ctypes.string_at(output.data, output.length)
    finally:
        kernel.LocalFree(output.data)


def session_path(tax_id: str) -> Path:
    if not os.environ.get("LOCALAPPDATA"):
        raise ExportError("LOCALAPPDATA is unavailable; disable session reuse.")
    return Path(os.environ["LOCALAPPDATA"]) / "GDTInvoiceExporter" / tax_id / "session.dpapi"


def save_session(path: Path, state: dict) -> None:
    protected = dpapi(json.dumps(state, ensure_ascii=False).encode("utf-8"))
    path.parent.mkdir(parents=True, exist_ok=True)
    temporary = path.with_suffix(".tmp")
    temporary.write_bytes(protected)
    temporary.replace(path)


def load_session(path: Path) -> dict:
    return json.loads(dpapi(path.read_bytes(), decrypt=True).decode("utf-8"))


def windows_keyring():
    try:
        import keyring
    except ImportError as exc:
        raise ExportError("Install keyring to use Windows Credential Manager.") from exc
    if os.name != "nt" or not type(keyring.get_keyring()).__module__.startswith("keyring.backends.Windows"):
        raise ExportError("Only the Windows Credential Manager backend is allowed.")
    return keyring
