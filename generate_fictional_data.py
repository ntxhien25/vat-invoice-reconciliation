"""Recreate only the explicitly fictional source fixtures. No client data required."""
from pathlib import Path
from datetime import datetime
import json,re
from openpyxl import Workbook
from openpyxl.styles import Font,PatternFill
ROOT=Path(__file__).resolve().parent
def generate():
    for spec in json.loads((ROOT/'fictional_inputs.json').read_text(encoding='utf-8')):
        wb=Workbook();wb.remove(wb.active)
        for sheet in spec['sheets']:
            ws=wb.create_sheet(sheet['name'])
            for row in sheet['rows']:
                ws.append([datetime.fromisoformat(v) if isinstance(v,str) and re.fullmatch(r'2025-\d{2}-\d{2}',v) else v for v in row])
            for row in ws:
                for c in row:
                    if isinstance(c.value,datetime):c.number_format='dd/mm/yyyy'
            for c in ws[sheet['header']]:c.font=Font(color='FFFFFF',bold=True);c.fill=PatternFill('solid',fgColor='16324F')
            ws.freeze_panes=f"A{sheet['header']+1}"
        wb.save(ROOT/spec['file'])
if __name__=='__main__':generate()
